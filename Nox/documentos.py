"""
Documentos de texto con formato real (.docx), pensados para dictado.

El caso de uso: Brian dicta ideas sueltas y Nox las convierte en un documento
ordenado, con titulos, listas y negritas. No es un editor: es "convertime este
chorro de ideas en algo presentable y abrimelo".

Se genera .docx porque abre en LibreOffice, en Word y en Google Docs: es el formato
con formato que mas lugares entienden.

El contenido llega en markdown, que es lo que el modelo escribe naturalmente, y aca
se traduce a formato de Word de verdad.
"""
import re
import sys
from datetime import datetime
from pathlib import Path

import config
import plataforma
import seguridad
from seguridad import Rechazado

if sys.stdout is not None:
    sys.stdout.reconfigure(encoding="utf-8")

EXTENSIONES = {".docx", ".txt", ".md", ".odt", ".pdf"}

MESES = (
    "enero", "febrero", "marzo", "abril", "mayo", "junio",
    "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre",
)


def _nombre_seguro(titulo: str) -> str:
    """
    Convierte un titulo en un nombre de archivo valido en cualquier sistema. Linux solo
    prohibe '/', pero Windows no acepta varios mas: se sacan todos, asi el archivo se
    puede copiar a un pendrive o mandar a alguien sin problemas.
    """
    limpio = re.sub(r'[<>:"/\\|?*]', "", titulo).strip(" .")
    limpio = re.sub(r"\s+", " ", limpio)
    return (limpio or "Documento")[:80]


def _aplicar_formato(parrafo, texto: str):
    """
    Pasa el markdown en linea (negritas y codigo) a formato de Word.

    Se parte el texto por los marcadores y se va agregando cada tramo con su estilo,
    porque Word no entiende asteriscos: hay que marcar cada fragmento.
    """
    for tramo in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", texto):
        if not tramo:
            continue
        if tramo.startswith("**") and tramo.endswith("**"):
            parrafo.add_run(tramo[2:-2]).bold = True
        elif tramo.startswith("`") and tramo.endswith("`"):
            r = parrafo.add_run(tramo[1:-1])
            r.font.name = "Consolas"
        elif tramo.startswith("*") and tramo.endswith("*") and len(tramo) > 2:
            parrafo.add_run(tramo[1:-1]).italic = True
        else:
            parrafo.add_run(tramo)


def crear_docx(titulo: str, contenido: str, ruta: Path) -> Path:
    """Escribe un .docx con el markdown ya traducido a formato de Word."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()

    # Cuerpo legible por defecto; Word trae 11pt Calibri, que esta bien.
    estilo = doc.styles["Normal"]
    estilo.font.size = Pt(11)

    doc.add_heading(titulo, level=0)
    # strftime("%B") devuelve el mes en el idioma del sistema (o en ingles), asi que
    # se arma a mano para que salga siempre en español.
    ahora = datetime.now()
    fecha = doc.add_paragraph(f"{ahora.day} de {MESES[ahora.month - 1]} de {ahora.year}")
    fecha.runs[0].italic = True

    for linea in contenido.split("\n"):
        limpia = linea.rstrip()

        if not limpia.strip():
            continue

        # Encabezados
        m = re.match(r"^(#{1,4})\s+(.*)$", limpia)
        if m:
            doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4))
            continue

        # Lista con viñetas
        m = re.match(r"^\s*[-*+]\s+(.*)$", limpia)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            _aplicar_formato(p, m.group(1))
            continue

        # Lista numerada
        m = re.match(r"^\s*\d+[.)]\s+(.*)$", limpia)
        if m:
            p = doc.add_paragraph(style="List Number")
            _aplicar_formato(p, m.group(1))
            continue

        # Cita
        if limpia.startswith(">"):
            p = doc.add_paragraph(style="Intense Quote")
            _aplicar_formato(p, limpia.lstrip("> ").strip())
            continue

        # Separador
        if re.fullmatch(r"[-*_]{3,}", limpia.strip()):
            doc.add_page_break()
            continue

        p = doc.add_paragraph()
        _aplicar_formato(p, limpia)

    doc.save(str(ruta))
    return ruta


def crear_documento(
    titulo: str, contenido: str, formato: str = "docx", abrir: bool = True
) -> str:
    """
    Crea un documento y opcionalmente lo abre.

    formato: 'docx' (con formato, abre en Word/LibreOffice), 'txt' o 'md' (texto plano).
    """
    formato = (formato or "docx").lower().lstrip(".")
    if f".{formato}" not in EXTENSIONES:
        return f"No manejo el formato '{formato}'. Uso: docx, txt o md."

    carpeta = config.CARPETA_DOCUMENTOS
    carpeta.mkdir(parents=True, exist_ok=True)
    destino = carpeta / f"{_nombre_seguro(titulo)}.{formato}"

    try:
        seguridad.validar_escritura(
            destino, [carpeta] + list(config.CARPETAS_ESCRITURA),
            config.EXTENSIONES_ESCRITURA | {".docx"},
        )
    except Rechazado as e:
        return f"Error: {e}"

    # Si ya existe, no pisarlo: se numera.
    if destino.exists():
        base, i = destino.stem, 2
        while destino.exists():
            destino = carpeta / f"{base} ({i}).{formato}"
            i += 1

    try:
        if formato == "docx":
            crear_docx(titulo, contenido, destino)
        else:
            encabezado = f"# {titulo}\n\n" if formato == "md" else f"{titulo}\n{'=' * len(titulo)}\n\n"
            destino.write_text(encabezado + contenido, encoding="utf-8")
    except Exception as e:
        return f"No pude crear el documento: {e}"

    mensaje = f"Documento creado: {destino}"

    if abrir:
        try:
            # Con el programa que el sistema tenga asociado a ese tipo de archivo
            # (Word o LibreOffice para .docx, un editor para .md...).
            plataforma.abrir_con_predeterminado(destino)
            mensaje += "\nLo abri para que lo revises."
        except OSError as e:
            mensaje += f"\n(No pude abrirlo solo: {e})"

    return mensaje


def ruta_libreoffice() -> Path | None:
    """Busca soffice, el conversor de LibreOffice por linea de comandos."""
    import os
    import shutil

    for nombre in ("soffice", "libreoffice"):
        ruta = shutil.which(nombre)
        if ruta:
            return Path(ruta)
    if plataforma.ES_WINDOWS:
        # El instalador de Windows no lo agrega al PATH.
        for base in (os.environ.get("ProgramFiles"), os.environ.get("ProgramFiles(x86)")):
            if base:
                candidata = Path(base) / "LibreOffice" / "program" / "soffice.exe"
                if candidata.exists():
                    return candidata
    return None


def convertir(ruta: str, a_formato: str = "pdf") -> str:
    """
    Convierte un documento a otro formato usando LibreOffice.

    Sirve para pasar a PDF (lo mas util: mandar algo sin que lo puedan editar), o a
    .odt / .docx / .rtf / .html segun haga falta.
    """
    import subprocess

    soffice = ruta_libreoffice()
    if soffice is None:
        return (
            "No encuentro LibreOffice instalado, que es lo que hace la conversion. "
            + ("Se instala con: winget install TheDocumentFoundation.LibreOffice"
               if plataforma.ES_WINDOWS else
               "Se instala con el gestor de paquetes (en Arch: sudo pacman -S libreoffice-fresh)")
        )

    origen = Path(ruta).expanduser()
    if not origen.is_absolute():
        origen = config.CARPETA_DOCUMENTOS / origen
    origen = seguridad.resolver(origen)

    if not origen.exists():
        return f"No existe: {origen}"

    formato = a_formato.lower().lstrip(".")
    if formato not in ("pdf", "odt", "docx", "rtf", "html", "txt"):
        return f"No convierto a '{formato}'. Uso: pdf, odt, docx, rtf, html o txt."

    try:
        r = subprocess.run(
            [str(soffice), "--headless", "--convert-to", formato,
             "--outdir", str(origen.parent), str(origen)],
            capture_output=True, text=True, timeout=180,
            encoding="utf-8", errors="ignore", **plataforma.kwargs_sin_ventana(),
        )
    except (OSError, subprocess.SubprocessError) as e:
        return f"Fallo la conversion: {e}"

    destino = origen.with_suffix(f".{formato}")
    if destino.exists():
        return f"Convertido: {destino}"
    return f"LibreOffice no genero el archivo. Salida: {(r.stdout or r.stderr)[:200]}"


def listar_documentos() -> str:
    """Documentos que Nox creo hasta ahora."""
    carpeta = config.CARPETA_DOCUMENTOS
    if not carpeta.exists():
        return "Todavia no cree ningun documento."

    archivos = sorted(
        [p for p in carpeta.iterdir() if p.suffix.lower() in EXTENSIONES],
        key=lambda p: -p.stat().st_mtime,
    )
    if not archivos:
        return "Todavia no cree ningun documento."

    lineas = [f"En {carpeta}:\n"]
    for p in archivos[:30]:
        fecha = datetime.fromtimestamp(p.stat().st_mtime).strftime("%d/%m/%Y %H:%M")
        kb = p.stat().st_size / 1024
        lineas.append(f"  {p.name}  ({kb:.0f} KB, {fecha})")
    return "\n".join(lineas)
